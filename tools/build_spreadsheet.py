"""Fill the ECON3700 unsealed roads workbook by editing sheet XML in place
(keeps every style, column width, freeze pane and web extension untouched)."""
import re, shutil, subprocess, sys, zipfile, os
from openpyxl.utils import get_column_letter

SRC, OUT, WORK = sys.argv[1], sys.argv[2], sys.argv[3]
YEARS = [get_column_letter(c) for c in range(3, 33)]  # C..AF
assert YEARS[0] == 'C' and YEARS[-1] == 'AF'

inputs, dcf = {}, {}

# ---------- Inputs sheet: avoided abatement cost ----------
inputs['C40'] = '1624'                        # Reef Trust Cape York gully benchmark, $/t/yr
inputs['C41'] = '(1+2.1%)*(1+3.8%)'           # ABS CPI: yr to Jun-25 2.1%, yr to Jun-26 3.8%
inputs['C42'] = 'C40*C41'

# ---------- DCF key parameters ----------
dcf['B5'] = 'Inputs!C2'
dcf['B6'] = 'Inputs!C3'
dcf['B7'] = 'Inputs!C8'
dcf['B8'] = 'Inputs!C9*Inputs!C12'
dcf['B9'] = 'Inputs!C12'
dcf['B10'] = 'Inputs!C42'
dcf['B11'] = 'Inputs!C22/100'

pv_rows = [17, 18, 19, 21, 22, 25, 26, 27, 28, 29, 31, 32,
           35, 36, 37, 38, 39, 40, 41, 42, 44, 45,
           48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59, 61, 62]
for r in pv_rows:
    dcf[f'B{r}'] = f'SUMPRODUCT(C{r}:AF{r},$C$14:$AF$14)'

def t(col):        # years since start (0 in FY2026)
    return f'({col}$13-Inputs!$C$4)'

grader = ('IF({c}$13-Inputs!$C$4<2,Inputs!$C$25,Inputs!$C$26)*$B$7')
weeds = ('CHOOSE(MIN({c}$13-Inputs!$C$4+1,4),Inputs!$C$28,Inputs!$C$29,'
         'Inputs!$C$30,Inputs!$C$31)*$B$7')

for c in YEARS:
    dcf[f'{c}14'] = f'1/(1+$B$5)^{t(c)}'
    # Scenario 1: business as usual
    dcf[f'{c}17'] = 'Inputs!$C$20*Inputs!$C$6'
    dcf[f'{c}18'] = 'Inputs!$C$21*Inputs!$C$7'
    dcf[f'{c}19'] = f'SUM({c}17:{c}18)'
    dcf[f'{c}20'] = 'Inputs!$C$35*$B$7'
    dcf[f'{c}21'] = f'{c}20*$B$10'
    dcf[f'{c}22'] = f'{c}19+{c}21'
    # Scenario 2: alternative maintenance
    dcf[f'{c}25'] = f'{c}17*$B$11'
    dcf[f'{c}26'] = f'{c}18*$B$11'
    dcf[f'{c}27'] = grader.format(c=c)
    dcf[f'{c}28'] = weeds.format(c=c)
    dcf[f'{c}29'] = f'SUM({c}25:{c}28)'
    dcf[f'{c}30'] = 'Inputs!$C$36*$B$7'
    dcf[f'{c}31'] = f'{c}30*$B$10'
    dcf[f'{c}32'] = f'{c}29+{c}31'
    # Scenario 3: major erosion control (capital in first year)
    first = f'IF({c}$13=Inputs!$C$4,1,0)'
    dcf[f'{c}35'] = f'{first}*Inputs!$C$14*$B$8'
    dcf[f'{c}36'] = f'{first}*Inputs!$C$15*$B$8'
    dcf[f'{c}37'] = f'{first}*Inputs!$C$16*$B$8'
    dcf[f'{c}38'] = f'{c}17*Inputs!$C$23/100'
    dcf[f'{c}39'] = f'{c}18*Inputs!$C$23/100'
    dcf[f'{c}40'] = grader.format(c=c)
    dcf[f'{c}41'] = weeds.format(c=c)
    dcf[f'{c}42'] = f'SUM({c}35:{c}41)'
    dcf[f'{c}43'] = 'Inputs!$C$37*$B$7'
    dcf[f'{c}44'] = f'{c}43*$B$10'
    dcf[f'{c}45'] = f'{c}42+{c}44'
    # Scenario 4: full betterment (bitumen dust seal over whole length)
    dcf[f'{c}48'] = f'{first}*Inputs!$C$14*$B$8'
    dcf[f'{c}49'] = f'{first}*Inputs!$C$15*$B$8'
    dcf[f'{c}50'] = f'{first}*Inputs!$C$16*$B$8'
    dcf[f'{c}51'] = f'{first}*Inputs!$C$17*$B$7'
    dcf[f'{c}52'] = f'{first}*Inputs!$C$18*$B$7'
    dcf[f'{c}53'] = f'{first}*{c}17'
    dcf[f'{c}54'] = f'{first}*{c}18'
    dcf[f'{c}55'] = f'(1-{first})*Inputs!$C$32*$B$7'
    dcf[f'{c}56'] = f'(1-{first})*Inputs!$C$33*$B$7'
    dcf[f'{c}57'] = grader.format(c=c)
    dcf[f'{c}58'] = weeds.format(c=c)
    dcf[f'{c}59'] = f'SUM({c}48:{c}58)'
    dcf[f'{c}60'] = 'Inputs!$C$38*$B$7'
    dcf[f'{c}61'] = f'{c}60*$B$10'
    dcf[f'{c}62'] = f'{c}59+{c}61'

# ---------- Results summary ----------
fin = {'B': 19, 'C': 29, 'D': 42, 'E': 59}
tot = {'B': 22, 'C': 32, 'D': 45, 'E': 62}
env = {'B': 21, 'C': 31, 'D': 44, 'E': 61}
sed = {'B': 20, 'C': 30, 'D': 43, 'E': 60}
for col in 'BCDE':
    dcf[f'{col}67'] = f'$B${fin[col]}'
    dcf[f'{col}68'] = f'$B${tot[col]}'
for col in 'CDE':
    dcf[f'{col}71'] = f'{col}67-$B$67'
    dcf[f'{col}72'] = f'$B${env["B"]}-$B${env[col]}'
    dcf[f'{col}74'] = f'-{col}71'
    dcf[f'{col}75'] = f'$B$68-{col}68'
    dcf[f'{col}78'] = f'SUM($C${sed["B"]}:$AF${sed["B"]})-SUM($C${sed[col]}:$AF${sed[col]})'
    dcf[f'{col}79'] = f'{col}71/{col}78'


def inject(xml, cells, values):
    for ref, formula in cells.items():
        col, row = re.match(r'([A-Z]+)(\d+)', ref).groups()
        v = values.get(ref) if values else None
        vtag = f'<v>{repr(float(v))}</v>' if v is not None else ''
        f = formula.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        new_inner = f'<f>{f}</f>{vtag}'
        pat = re.compile(r'<c r="%s"( s="\d+")?(?: t="\w+")?\s*(?:/>|>.*?</c>)' % ref, re.S)
        m = pat.search(xml)
        if not m:
            raise SystemExit(f'cell {ref} not present in template')
        xml = xml[:m.start()] + f'<c r="{ref}"{m.group(1) or ""}>{new_inner}</c>' + xml[m.end():]
    return xml


def build(dst, values=None):
    zin = zipfile.ZipFile(SRC)
    zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        name = item.filename
        if name == 'xl/calcChain.xml':
            continue  # Excel rebuilds it
        if name == '[Content_Types].xml':
            data = re.sub(rb'<Override PartName="/xl/calcChain.xml"[^>]*/>', b'', data)
        if name == 'xl/_rels/workbook.xml.rels':
            data = re.sub(rb'<Relationship [^>]*Target="calcChain.xml"/>', b'', data)
        if name == 'xl/workbook.xml':
            s = data.decode()
            s = re.sub(r'<calcPr ([^/]*)/>', lambda m: '<calcPr ' + re.sub(r'\s*fullCalcOnLoad="\d"', '', m.group(1)) + ' fullCalcOnLoad="1"/>', s)
            data = s.encode()
        if name == 'xl/worksheets/sheet1.xml':
            data = inject(data.decode(), inputs, (values or {}).get('Inputs')).encode()
        if name == 'xl/worksheets/sheet2.xml':
            data = inject(data.decode(), dcf, (values or {}).get('DCF')).encode()
        zout.writestr(item, data)
    zout.close()


os.makedirs(WORK, exist_ok=True)
stage1 = os.path.join(WORK, 'stage1.xlsx')
build(stage1)
# Recalculate with LibreOffice to obtain cached values, then write them into the untouched-format file.
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir',
                os.path.join(WORK, 'lo'), stage1], check=True, capture_output=True)
import openpyxl
wb = openpyxl.load_workbook(os.path.join(WORK, 'lo', 'stage1.xlsx'), data_only=True)
values = {'Inputs': {k: wb['Inputs'][k].value for k in inputs},
          'DCF': {k: wb['DCF'][k].value for k in dcf}}
missing = [k for s in values for k, v in values[s].items() if not isinstance(v, (int, float))]
if missing:
    raise SystemExit(f'non-numeric results: {missing[:10]}')
build(OUT, values)
print('ok', OUT)
